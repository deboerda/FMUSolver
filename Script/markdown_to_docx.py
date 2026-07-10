import os
import sys
import re
try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
except ImportError:
    print("[!] python-docx is missing. Please run: pip install python-docx")
    sys.exit(1)

def set_cell_background(cell, fill_hex):
    """设置 Word 单元格背景颜色"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """设置 Word 单元格内边距 (Padding)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        m = OxmlElement(margin)
        m.set(qn('w:w'), str(val))
        m.set(qn('w:type'), 'dxa')
        tcMar.append(m)
    tcPr.append(tcMar)

def parse_inline(paragraph, text):
    """解析 Markdown 行内样式：***粗斜体***, **粗体**, *斜体*, `代码`"""
    pattern = re.compile(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`)')
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('***') and part.endswith('***'):
            run = paragraph.add_run(part[3:-3])
            run.bold = True
            run.italic = True
        elif part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(165, 42, 42)  # 深红褐色标识行内代码
        else:
            paragraph.add_run(part)

def parse_and_add_table(doc, table_rows):
    """解析 Markdown 表格行并添加到 Word 中"""
    raw_cells_rows = []
    for r in table_rows:
        row_cells = [c.strip() for c in r.split('|')]
        # 去除头尾分割产生的空字符串
        if len(row_cells) > 1 and row_cells[0] == '':
            row_cells = row_cells[1:]
        if len(row_cells) > 0 and row_cells[-1] == '':
            row_cells = row_cells[:-1]
        raw_cells_rows.append(row_cells)
        
    if not raw_cells_rows:
        return
        
    # 剔除表头分割线行 |---|---|
    if len(raw_cells_rows) > 1:
        second_row = raw_cells_rows[1]
        if all(re.match(r'^:?\-+:?$', cell) for cell in second_row):
            del raw_cells_rows[1]
            
    num_rows = len(raw_cells_rows)
    num_cols = max(len(r) for r in raw_cells_rows)
    
    table = doc.add_table(rows=num_rows, cols=num_cols)
    table.autofit = True
    
    # 插入轻量级表格边框 (仅横向浅灰框，隐藏纵向框，符合现代排版)
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)
    
    for r_idx, row_cells in enumerate(raw_cells_rows):
        row = table.rows[r_idx]
        for c_idx, cell_text in enumerate(row_cells):
            if c_idx < num_cols:
                cell = row.cells[c_idx]
                cell.text = ""  # 清空默认文本占位
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                
                if r_idx == 0:
                    # 表头行：加粗，白字，深绿色背景
                    run = p.add_run(cell_text)
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    set_cell_background(cell, "1B5E20")
                else:
                    parse_inline(p, cell_text)
                    # 交替行着色 (Alternating Row Color)
                    if r_idx % 2 == 1:
                        set_cell_background(cell, "F9F9F9")
                
                # 设定内边距
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

def markdown_to_docx(md_path, docx_path):
    """主转换逻辑：从 Markdown 转换为 Word DOCX 格式"""
    print(f"[*] Converting: {os.path.basename(md_path)} -> {os.path.basename(docx_path)}")
    doc = Document()
    
    # 基础文档字体配置
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Microsoft YaHei'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(51, 51, 51)
    
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()
        
    lines = content.split('\n')
    
    in_code_block = False
    code_lines = []
    in_table = False
    table_rows = []
    
    i = 0
    while i < len(lines):
        line = lines[i]
        
        # --- 代码块处理 ---
        if line.strip().startswith('```'):
            if in_code_block:
                in_code_block = False
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(6)
                
                # 左侧代码高亮竖线及浅灰背景
                pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="8" w:color="CCCCCC"/></w:pBdr>')
                p._p.get_or_add_pPr().append(pBdr)
                shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5"/>')
                p._p.get_or_add_pPr().append(shd)
                
                run = p.add_run('\n'.join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(85, 85, 85)
                code_lines = []
            else:
                in_code_block = True
            i += 1
            continue
            
        if in_code_block:
            code_lines.append(line)
            i += 1
            continue
            
        # --- 表格处理 ---
        if line.strip().startswith('|'):
            in_table = True
            table_rows.append(line)
            i += 1
            continue
        elif in_table:
            # 退出表格区域，先渲染之前的表格
            parse_and_add_table(doc, table_rows)
            table_rows = []
            in_table = False
            # 不要 continue，继续往下处理当前行
            
        # --- 标题处理 (#) ---
        m_heading = re.match(r'^(#{1,6})\s+(.*)$', line.strip())
        if m_heading:
            level = len(m_heading.group(1))
            title = m_heading.group(2)
            heading = doc.add_heading(level=level)
            heading.paragraph_format.space_before = Pt(14)
            heading.paragraph_format.space_after = Pt(6)
            run = heading.add_run(title)
            run.bold = True
            run.font.name = 'Microsoft YaHei'
            
            # 分级配置颜色和字体大小
            if level == 1:
                run.font.size = Pt(18)
                run.font.color.rgb = RGBColor(13, 71, 161)  # 深蓝色
            elif level == 2:
                run.font.size = Pt(14)
                run.font.color.rgb = RGBColor(27, 94, 32)   # 深绿色
            else:
                run.font.size = Pt(11.5)
                run.font.color.rgb = RGBColor(69, 90, 100)  # 蓝灰色
            i += 1
            continue
            
        # --- 水平分割线 ---
        if line.strip() in ('---', '***', '___'):
            p = doc.add_paragraph()
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="DDDDDD"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue
            
        # --- 无序列表 (-, *, +) ---
        m_ul = re.match(r'^([\-\*\+])\s+(.*)$', line.strip())
        if m_ul:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            parse_inline(p, m_ul.group(2))
            i += 1
            continue
            
        # --- 有序列表 (\d+.) ---
        m_ol = re.match(r'^(\d+)\.\s+(.*)$', line.strip())
        if m_ol:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            parse_inline(p, m_ol.group(2))
            i += 1
            continue
            
        # --- 引用与警告框 (> [!NOTE] 等) ---
        m_bq = re.match(r'^>\s*(.*)$', line.strip())
        if m_bq:
            bq_content = m_bq.group(1).strip()
            # 兼容 GitHub Alerts 语法
            m_alert = re.match(r'^\[!(NOTE|TIP|IMPORTANT|WARNING|CAUTION)\]\s*(.*)$', bq_content)
            alert_type = None
            alert_text = bq_content
            if m_alert:
                alert_type = m_alert.group(1)
                alert_text = m_alert.group(2)
                
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            
            # 根据类型定制边框颜色和背景填充
            bdr_color = "CCCCCC"
            fill_color = "FAFAFA"
            if alert_type == "NOTE":
                bdr_color = "0288D1"  # 蓝色
                fill_color = "E1F5FE"
            elif alert_type == "TIP":
                bdr_color = "388E3C"  # 绿色
                fill_color = "E8F5E9"
            elif alert_type in ("WARNING", "IMPORTANT"):
                bdr_color = "F57C00"  # 橙色
                fill_color = "FFF3E0"
            elif alert_type == "CAUTION":
                bdr_color = "D32F2F"  # 红色
                fill_color = "FFEBEE"
                
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="8" w:color="{bdr_color}"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
            p._p.get_or_add_pPr().append(shd)
            
            if alert_type:
                run_prefix = p.add_run(f"【{alert_type}】")
                run_prefix.bold = True
                run_prefix.font.color.rgb = RGBColor(
                    int(bdr_color[0:2], 16),
                    int(bdr_color[2:4], 16),
                    int(bdr_color[4:6], 16)
                )
            
            parse_inline(p, alert_text)
            i += 1
            continue
            
            
        # 空行忽略
        if not line.strip():
            i += 1
            continue
            
        # --- 普通段落 ---
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        parse_inline(p, line.strip())
        i += 1
        
    # 强制在表格没有被包裹完退出文件时刷出最后一个表格
    if in_table and table_rows:
        parse_and_add_table(doc, table_rows)

    doc.save(docx_path)
    print(f"[OK] Conversion complete: {docx_path}")

if __name__ == "__main__":
    if len(sys.argv) > 1:
        md_file = sys.argv[1]
        if not os.path.exists(md_file):
            print(f"[!] File not found: {md_file}")
            sys.exit(1)
        docx_file = sys.argv[2] if len(sys.argv) > 2 else md_file.replace('.md', '.docx')
        markdown_to_docx(md_file, docx_file)
    else:
        # 默认模式：自动查找当前目录和父目录下的 FmuSolver使用操作和维护说明书.md
        possible_files = [
            "FmuSolver使用操作和维护说明书.md",
            "../FmuSolver使用操作和维护说明书.md",
            "../../FmuSolver使用操作和维护说明书.md"
        ]
        converted = False
        for f in possible_files:
            if os.path.exists(f):
                out = f.replace('.md', '.docx')
                markdown_to_docx(f, out)
                converted = True
                break
        if not converted:
            print("[*] No default manual file found. Usage: python markdown_to_docx.py <input.md> [output.docx]")
            # 扫描当前目录所有.md
            md_files = [f for f in os.listdir('.') if f.endswith('.md')]
            if md_files:
                print(f"[*] Found {len(md_files)} markdown files in current directory. Converting all...")
                for f in md_files:
                    markdown_to_docx(f, f.replace('.md', '.docx'))
            else:
                print("[!] No markdown files found.")
